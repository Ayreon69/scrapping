"""
Charge des profils MRH depuis profils_template.xlsx.
Convertit les libelles lisibles en valeurs API.
"""
import random
from datetime import date, timedelta
from faker import Faker
from openpyxl import load_workbook

fake = Faker("fr_FR")

# ---------------------------------------------------------------------------
# Tables de conversion libelle -> valeur API
# ---------------------------------------------------------------------------

STATUT_MAP = {
    "Locataire":                          "LOCATAIRE_COLOCATAIRE",
    "Locataire meuble":                   "LOCATAIRE_COLOCATAIRE_MEUBLE",
    "Proprietaire occupant":              "PROPRIETAIRE_OCCUPANT",
    "Proprietaire non occupant (PNO)":    "PROPRIETAIRE_NON_OCCUPANT",
}

HABITATION_MAP = {
    "Appartement": "APPARTEMENT",
    "Maison":      "MAISON",
}

RESIDENCE_MAP = {
    "Residence principale":  "RESIDENCE_PRINCIPALE",
    "Residence secondaire":  "RESIDENCE_SECONDAIRE",
}

ANCIENNETE_MAP = {
    "Moins de 5 ans":       "MOINS_DE_5_ANS",
    "Entre 5 et 10 ans":    "ENTRE_5_ET_10_ANS",
    "Entre 10 et 30 ans":   "ENTRE_10_ET_30_ANS",
    "Plus de 30 ans":       "PLUS_DE_30_ANS",
}

ETAGE_MAP = {
    "Rez-de-chaussee": "REZ_DE_CHAUSSEE",
    "Autre etage":     "AUTRE_ETG",
    "Dernier etage":   "DERNIER_ETG",
}

JOURS_MAP = {
    "Moins de 45 jours":      "MOINS_45_JRS",
    "Entre 45 et 60 jours":   "ENTRE_45_ET_60_JRS",
    "Entre 60 et 90 jours":   "ENTRE_60_ET_90_JRS",
    "Plus de 90 jours":       "PLUS_DE_90_JRS",
    "Plus de 180 jours":      "PLUS_DE_180_JRS",
}


CHAUFFAGE_MAP = {
    "Poele a bois":          "poele_a_bois",
    "Cheminee sans insert":  "cheminee_sans_insert",
    "Cheminee avec insert":  "cheminee_insert",
}

CHEMINEE_PRO_MAP = {
    "Oui":          "OUI",
    "Non":          "NON",
    "Ne sait pas":  "NSP",
}

PISCINE_MAP = {
    "Piscine interieure":             "piscine_interieure",
    "Piscine exterieure couverte":    "piscine_exterieure_couverte",
    "Piscine exterieure non couverte":"piscine_exterieure_non_couverte",
}

ETAT_ASSUREUR_MAP = {
    "Oui, je suis assure":           "OUI",
    "Non, je n ai pas d assurance":  "NON",
    "Je n ai jamais ete assure":     "jamais-assure",
}

ASSUREUR_MAP = {
    "AXA": "AXA", "Allianz": "Allianz", "Eurofil": "Eurofil",
    "Thelem": "Thelem", "Credit Agricole": "Credit Agricole",
    "Banque Populaire": "Banque Populaire", "Generali": "Generali",
    "GMF": "GMF", "Groupama": "Groupama", "MAAF": "MAAF",
    "Matmut": "Matmut", "MAIF": "MAIF",
    "Aucun de ces assureurs": "",
}

PROFESSION_MAP = {
    "Salarie":                      "salarie",
    "Salarie cadre":                "salarie_cadre",
    "Retraite":                     "retraite",
    "Etudiant":                     "etudiant",
    "Fonction publique etat":       "Fonction_publique_d_etat",
    "Fonction publique territoriale": "Fonction_publique_territoriale",
    "Fonction publique hospitaliere": "Fonction_publique_hospitaliere",
    "Artisan":                      "artisan",
    "Commercant":                   "commercant",
    "Profession liberale":          "profession_liberale",
    "Chef d entreprise":            "chef_d_entreprise",
    "Enseignant":                   "enseignant",
    "Agriculteur":                  "agriculteur",
    "Exploitant agricole":          "exploitant_agricole",
    "VRP":                          "vrp",
    "Visiteur medical":             "visiteur_medical",
    "En recherche d emploi":        "en_recherche_d_emploi",
    "Sans profession":              "sans_profession",
}

SITUATION_MAP = {
    "Celibataire":      "celibataire",
    "Marie(e)":         "marie",
    "Divorce(e)":       "divorce",
    "Pacse(e)":         "pacse",
    "En concubinage":   "en-concubinage",
    "Separe(e)":        "separe",
    "Veuf(ve)":         "veuf",
}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _parse_date(val) -> str:
    s = _str(val)
    if not s:
        return ""
    if "/" in s:
        parts = s.split("/")
        if len(parts) == 3:
            return f"{parts[2]}-{parts[1].zfill(2)}-{parts[0].zfill(2)}"
    return s


def _int_or(val, default=0) -> int:
    try:
        return int(float(str(val)))
    except (TypeError, ValueError):
        return default


def _random_phone() -> str:
    return "0" + str(random.randint(600000000, 799999999))


def random_pieces(surface: int) -> dict:
    s = int(surface)
    if s < 30:
        return {"inferieures_a_30": 1, "entre_30_40": 0, "entre_40_50": 0, "superieures_50": 0}
    elif s < 50:
        return {"inferieures_a_30": 1, "entre_30_40": 1, "entre_40_50": 0, "superieures_50": 0}
    elif s < 70:
        return {"inferieures_a_30": 1, "entre_30_40": 1, "entre_40_50": 1, "superieures_50": 0}
    else:
        nb_sup = (s - 60) // 20
        return {"inferieures_a_30": 1, "entre_30_40": 1, "entre_40_50": 1, "superieures_50": nb_sup}


# ---------------------------------------------------------------------------
# Fonction principale
# ---------------------------------------------------------------------------

def load_profiles_from_excel(path: str) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    if "Profils" not in wb.sheetnames:
        raise ValueError(f"Feuille 'Profils' introuvable dans {path}")

    ws = wb["Profils"]
    headers = [_str(cell.value) for cell in ws[2]]

    date_debut_contact = (date.today() + timedelta(days=1)).strftime("%Y-%m-%d")
    annee_courante = date.today().year
    profiles = []

    for row in ws.iter_rows(min_row=3, values_only=True):
        if all(v is None or _str(v) == "" for v in row):
            continue

        r = {headers[i]: _str(v) for i, v in enumerate(row) if i < len(headers)}

        # Logement
        statut       = STATUT_MAP.get(r.get("Statut occupant", ""), "LOCATAIRE_COLOCATAIRE")
        type_hab     = HABITATION_MAP.get(r.get("Type de bien", "Appartement"), "APPARTEMENT")
        type_res     = RESIDENCE_MAP.get(r.get("Type de residence", ""), "RESIDENCE_PRINCIPALE")
        anciennete   = ANCIENNETE_MAP.get(r.get("Anciennete logement", ""), "ENTRE_5_ET_10_ANS")
        jours        = JOURS_MAP.get(r.get("Jours inhabite/an", ""), "MOINS_45_JRS")

        etage_label  = r.get("Etage", "")
        # Pour une maison, l'étage n'est pas applicable — forcer AUTRE_ETG
        etage        = "AUTRE_ETG" if type_hab == "MAISON" else ETAGE_MAP.get(etage_label, "AUTRE_ETG")

        surface      = _int_or(r.get("Surface (m2)", 60), 60)
        surface      = surface if surface > 0 else 60

        # Equipements : une colonne OUI/NON par equipement
        equipements = []
        if r.get("Garage", "NON").upper() == "OUI":
            equipements.append("Garage")
        if r.get("Systeme d alarme", "NON").upper() == "OUI":
            equipements.append("antivol")
        if r.get("Veranda / Loggia", "NON").upper() == "OUI":
            equipements.append("Veranda")
        if r.get("Annexes / Dependances", "NON").upper() == "OUI":
            equipements.append("Annexes")
        if r.get("Chauffage bois/cheminee", "NON").upper() == "OUI":
            equipements.append("chauffage")
        if r.get("Piscine enterree", "NON").upper() == "OUI":
            equipements.append("piscine")

        chauffage_label = r.get("Type chauffage", "")
        type_chauffage  = CHAUFFAGE_MAP.get(chauffage_label, "") if "chauffage" in equipements else ""

        cheminee_pro_label = r.get("Cheminee entretenue pro ?", "")
        cheminee_pro = CHEMINEE_PRO_MAP.get(cheminee_pro_label, "") if type_chauffage in ("cheminee_sans_insert", "cheminee_insert") else ""

        piscine_label = r.get("Type piscine", "")
        type_piscine  = PISCINE_MAP.get(piscine_label, "") if "piscine" in equipements else ""

        surface_veranda = r.get("Surface veranda (m2)", "") if "Veranda" in equipements else ""
        surface_annexes = r.get("Surface annexes (m2)", "") if "Annexes" in equipements else ""

        capital_mobilier = _int_or(r.get("Capital mobilier (EUR)", 15000), 15000)
        if capital_mobilier <= 0:
            capital_mobilier = 15000

        # Occupants
        nb_adultes = _int_or(r.get("Nb adultes", 1), 1)
        nb_adultes = max(1, min(2, nb_adultes))
        nb_enfants = _int_or(r.get("Nb enfants", 0), 0)
        nb_enfants = max(0, min(3, nb_enfants))

        enfants = {}
        for i in range(1, 10):
            enfants[f"date_naissance_enfant{i}"] = ""

        # Assurance actuelle
        etat_label   = r.get("Deja assure ?", "")
        etat_assureur = ETAT_ASSUREUR_MAP.get(etat_label, "NON")
        assureur_label = r.get("Assureur actuel", "")
        assureur_hab   = ASSUREUR_MAP.get(assureur_label, "") if etat_assureur == "OUI" else ""

        # Profil
        date_naissance = _parse_date(r.get("Date de naissance", ""))
        situation      = SITUATION_MAP.get(r.get("Situation maritale", ""), "celibataire")
        profession     = PROFESSION_MAP.get(r.get("Profession", "Salarie"), "salarie")

        # Identite auto
        civilite  = random.choice(["Monsieur", "Madame"])
        nom       = fake.last_name()
        prenom    = fake.first_name_male() if civilite == "Monsieur" else fake.first_name_female()
        email     = fake.email()
        telephone = _random_phone()

        cp    = r.get("Code postal", "") or "75001"
        ville = r.get("Ville", "")       or "Paris"

        annee_demenagement = str(random.randint(2010, annee_courante))

        test_label = " | ".join(filter(None, [
            r.get("Statut occupant", ""),
            r.get("Type de bien", ""),
            r.get("Profession", ""),
        ]))

        profile = {
            "id_forms": 3,
            "type_habitation":    type_hab,
            "etage_appartement":  etage,
            "type_residence":     type_res,
            "fins_professionnelles": "NON",
            "deja_proprietaire_logement": "",
            "logement_proposer_a_location": "",
            "duree_location": "",
            "emmenagement_prevus": "",
            "logement_renovation": "",
            "type_chauffage":     type_chauffage,
            "cheminee_professionel": cheminee_pro,
            "type_piscine":       type_piscine,
            "anciennete_logement": anciennete,
            "surface_habitable":  str(surface),
            **random_pieces(surface),
            "equipent_logement":  equipements,
            "surface_veranda":    str(surface_veranda),
            "surface_annexes":    str(surface_annexes),
            "nbr_jours_inhabita": jours,
            "moyen_de_protection": "NON",
            "moyen_de_protection_detail": [],
            "mean_protection_other": "",
            "statut_resident":    statut,
            "logement_cond":      "NON" if statut == "PROPRIETAIRE_NON_OCCUPANT" else "OUI",
            "annee_demmenagement": annee_demenagement,
            "nbr_adultes":        nb_adultes,
            "nbr_enfants":        nb_enfants,
            "capital_mobilier":   capital_mobilier,
            "res_3_dernieres_ann": "NON",
            "nbr_sinistres": 0,
            "type_sinistre1": "", "mois_SinDate1": "01", "annee_SinDate1": "",
            "type_sinistre2": "", "mois_SinDate2": "01", "annee_SinDate2": "",
            "type_sinistre3": "", "mois_SinDate3": "01", "annee_SinDate3": "",
            "type_sinistre4": "", "mois_SinDate4": "01", "annee_SinDate4": "",
            "date_sinistre1": "", "date_sinistre2": "", "date_sinistre3": "", "date_sinistre4": "",
            "etat_assureur":      etat_assureur,
            "assureur_habitation": assureur_hab,
            "date_debut_contact": date_debut_contact,
            "civilite":           civilite,
            "nom":                nom,
            "prenom":             prenom,
            "date_naissance":     date_naissance,
            **enfants,
            "adresse":            fake.street_address(),
            "cp":                 cp,
            "cp_logment":         cp,
            "ville":              ville,
            "ville_logment":      ville,
            "situation_matrimoniale": situation,
            "profession":         profession,
            "email":              email,
            "telephone":          telephone,
            "accepte_news":       False,
            "accepte_offre":      False,
            "provenance":         None,
            "_test_label":        test_label,
        }
        profiles.append(profile)

    if not profiles:
        raise ValueError(f"Aucun profil trouve dans {path}")

    return profiles
