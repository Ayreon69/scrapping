from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

wb = Workbook()
ws = wb.active
ws.title = "Profils"
ws_lists = wb.create_sheet("Listes")

# ---------------------------------------------------------------------------
# Listes deroulantes
# ---------------------------------------------------------------------------
LISTS = {
    "Statut":              ["Locataire", "Locataire meuble", "Proprietaire occupant", "Proprietaire non occupant (PNO)"],
    "Type habitation":     ["Appartement", "Maison"],
    "Type residence":      ["Residence principale", "Residence secondaire"],
    "Anciennete":          ["Moins de 5 ans", "Entre 5 et 10 ans", "Entre 10 et 30 ans", "Plus de 30 ans"],
    "Etage":               ["Rez-de-chaussee", "Autre etage", "Dernier etage"],
    "Jours inhabite":      ["Moins de 45 jours", "Entre 45 et 60 jours", "Entre 60 et 90 jours", "Plus de 90 jours", "Plus de 180 jours"],
    "Oui Non":             ["OUI", "NON"],
    "Chauffage":           ["Poele a bois", "Cheminee sans insert", "Cheminee avec insert"],
    "Cheminee pro":        ["Oui", "Non", "Ne sait pas"],
    "Piscine":             ["Piscine interieure", "Piscine exterieure couverte", "Piscine exterieure non couverte"],
    "Assureur actuel":     ["AXA", "Allianz", "Eurofil", "Thelem", "Credit Agricole", "Banque Populaire", "Generali", "GMF", "Groupama", "MAAF", "Matmut", "MAIF", "Aucun de ces assureurs"],
    "Etat assureur":       ["Oui, je suis assure", "Non, je n ai pas d assurance", "Je n ai jamais ete assure"],
    "Profession":          ["Salarie", "Salarie cadre", "Retraite", "Etudiant", "Fonction publique etat", "Fonction publique territoriale", "Fonction publique hospitaliere", "Artisan", "Commercant", "Profession liberale", "Chef d entreprise", "Enseignant", "Agriculteur", "Exploitant agricole", "VRP", "Visiteur medical", "En recherche d emploi", "Sans profession"],
    "Situation maritale":  ["Celibataire", "Marie(e)", "Divorce(e)", "Pacse(e)", "En concubinage", "Separe(e)", "Veuf(ve)"],
}

list_col_map = {}
for col_idx, (field, values) in enumerate(LISTS.items(), start=1):
    col_letter = get_column_letter(col_idx)
    list_col_map[field] = (col_letter, len(values))
    ws_lists.cell(row=1, column=col_idx, value=field).font = Font(bold=True)
    for row_idx, val in enumerate(values, start=2):
        ws_lists.cell(row=row_idx, column=col_idx, value=val)

# ---------------------------------------------------------------------------
# Colonnes : (header, cle_interne, list_key_ou_None, commentaire)
# ---------------------------------------------------------------------------
COLUMNS = [
    # Logement
    ("Statut occupant",         "statut",           "Statut",           "Locataire / Proprietaire occupant / PNO..."),
    ("Type de bien",            "type_habitation",  "Type habitation",  "Appartement ou Maison"),
    ("Type de residence",       "type_residence",   "Type residence",   "Principale ou Secondaire"),
    ("Etage",                   "etage",            "Etage",            "Uniquement pour un appartement - laisser vide pour une maison"),
    ("Surface (m2)",            "surface",          None,               "Surface habitable en m2 - ex: 60"),
    ("Anciennete logement",     "anciennete",       "Anciennete",       "Depuis combien de temps etes-vous dans ce logement"),
    ("Jours inhabite/an",       "jours_inhabita",   "Jours inhabite",   "Combien de jours par an le logement est inoccupe"),

    # Equipements : une colonne OUI/NON par equipement
    ("Garage",                  "eq_garage",        "Oui Non",          "Le logement dispose-t-il d un garage ?"),
    ("Systeme d alarme",        "eq_antivol",       "Oui Non",          "Le logement dispose-t-il d un systeme d alarme / antivol ?"),
    ("Veranda / Loggia",        "eq_veranda",       "Oui Non",          "Le logement dispose-t-il d une veranda ou loggia ?"),
    ("Surface veranda (m2)",    "surface_veranda",  None,               "Remplir uniquement si Veranda = OUI - ex: 15"),
    ("Annexes / Dependances",   "eq_annexes",       "Oui Non",          "Le logement dispose-t-il d annexes ou dependances ?"),
    ("Surface annexes (m2)",    "surface_annexes",  None,               "Remplir uniquement si Annexes = OUI - ex: 20"),
    ("Chauffage bois/cheminee", "eq_chauffage",     "Oui Non",          "Le logement dispose-t-il d un chauffage bois ou cheminee ?"),
    ("Type chauffage",          "chauffage",        "Chauffage",        "Remplir uniquement si Chauffage = OUI"),
    ("Cheminee entretenue pro ?","cheminee_pro",    "Cheminee pro",     "Remplir uniquement si cheminee avec ou sans insert"),
    ("Piscine enterree",        "eq_piscine",       "Oui Non",          "Le logement dispose-t-il d une piscine enterree ?"),
    ("Type piscine",            "piscine",          "Piscine",          "Remplir uniquement si Piscine = OUI"),

    # Occupants
    ("Capital mobilier (EUR)",  "capital_mobilier", None,               "Valeur estimee du mobilier - ex: 15000"),
    ("Nb adultes",              "nb_adultes",       None,               "Nombre d adultes dans le logement (1 ou 2)"),
    ("Nb enfants",              "nb_enfants",       None,               "Nombre d enfants dans le logement (0 a 3)"),

    # Assurance actuelle
    ("Deja assure ?",           "etat_assureur",    "Etat assureur",    "Avez-vous deja une assurance habitation ?"),
    ("Assureur actuel",         "assureur",         "Assureur actuel",  "Remplir uniquement si vous etes deja assure"),

    # Profil
    ("Date de naissance",       "date_naissance",   None,               "Format JJ/MM/AAAA - ex: 15/06/1985"),
    ("Situation maritale",      "situation",        "Situation maritale","Situation familiale"),
    ("Profession",              "profession",       "Profession",       "Activite professionnelle"),

    # Localisation
    ("Code postal",             "cp",               None,               "Ex: 75001 - laisser vide = Paris"),
    ("Ville",                   "ville",            None,               "Ex: Paris - laisser vide = Paris"),
]

# ---------------------------------------------------------------------------
# Couleurs par section
# ---------------------------------------------------------------------------
BLOC_COLORS = {
    "statut": "1A5276", "type_habitation": "1A5276", "type_residence": "1A5276",
    "etage": "1A5276", "surface": "1A5276", "anciennete": "1A5276", "jours_inhabita": "1A5276",
    "eq_garage": "1F618D", "eq_antivol": "1F618D", "eq_veranda": "1F618D",
    "surface_veranda": "1F618D", "eq_annexes": "1F618D", "surface_annexes": "1F618D",
    "eq_chauffage": "1F618D", "chauffage": "1F618D", "cheminee_pro": "1F618D",
    "eq_piscine": "1F618D", "piscine": "1F618D",
    "capital_mobilier": "117A65", "nb_adultes": "117A65", "nb_enfants": "117A65",
    "etat_assureur": "7D6608", "assureur": "7D6608",
    "date_naissance": "6E2F84", "situation": "6E2F84", "profession": "6E2F84",
    "cp": "784212", "ville": "784212",
}

BLOC_CELL_COLORS = {
    "statut": "EBF5FB", "type_habitation": "EBF5FB", "type_residence": "EBF5FB",
    "etage": "EBF5FB", "surface": "EBF5FB", "anciennete": "EBF5FB", "jours_inhabita": "EBF5FB",
    "eq_garage": "D6EAF8", "eq_antivol": "D6EAF8", "eq_veranda": "D6EAF8",
    "surface_veranda": "D6EAF8", "eq_annexes": "D6EAF8", "surface_annexes": "D6EAF8",
    "eq_chauffage": "D6EAF8", "chauffage": "D6EAF8", "cheminee_pro": "D6EAF8",
    "eq_piscine": "D6EAF8", "piscine": "D6EAF8",
    "capital_mobilier": "E8F8F5", "nb_adultes": "E8F8F5", "nb_enfants": "E8F8F5",
    "etat_assureur": "FEFBD8", "assureur": "FEFBD8",
    "date_naissance": "F5EEF8", "situation": "F5EEF8", "profession": "F5EEF8",
    "cp": "FAF0E6", "ville": "FAF0E6",
}

thin = Side(style="thin", color="BBBBBB")
border_normal = Border(left=thin, right=thin, top=thin, bottom=thin)
NB_ROWS = 50

# Calcul dynamique des plages de colonnes pour les blocs
def col_range(fields):
    indices = [i+1 for i, (_, f, _, _) in enumerate(COLUMNS) if f in fields]
    return get_column_letter(min(indices)), get_column_letter(max(indices))

LOGEMENT_FIELDS   = ["statut","type_habitation","type_residence","etage","surface","anciennete","jours_inhabita"]
EQUIP_FIELDS      = ["eq_garage","eq_antivol","eq_veranda","surface_veranda","eq_annexes","surface_annexes","eq_chauffage","chauffage","cheminee_pro","eq_piscine","piscine"]
OCCUPANTS_FIELDS  = ["capital_mobilier","nb_adultes","nb_enfants"]
ASSURANCE_FIELDS  = ["etat_assureur","assureur"]
PROFIL_FIELDS     = ["date_naissance","situation","profession"]
LOCAL_FIELDS      = ["cp","ville"]

c1s, c1e = col_range(LOGEMENT_FIELDS)
c2s, c2e = col_range(EQUIP_FIELDS)
c3s, c3e = col_range(OCCUPANTS_FIELDS)
c4s, c4e = col_range(ASSURANCE_FIELDS)
c5s, c5e = col_range(PROFIL_FIELDS)
c6s, c6e = col_range(LOCAL_FIELDS)

BLOC_LABELS = [
    (f"{c1s}1:{c1e}1", "LOGEMENT",           "1A5276"),
    (f"{c2s}1:{c2e}1", "EQUIPEMENTS",        "1F618D"),
    (f"{c3s}1:{c3e}1", "OCCUPANTS",          "117A65"),
    (f"{c4s}1:{c4e}1", "ASSURANCE ACTUELLE", "7D6608"),
    (f"{c5s}1:{c5e}1", "PROFIL",             "6E2F84"),
    (f"{c6s}1:{c6e}1", "LOCALISATION",       "784212"),
]

ws.row_dimensions[1].height = 22
for cell_range, label, color in BLOC_LABELS:
    ws.merge_cells(cell_range)
    first_cell = ws[cell_range.split(":")[0]]
    first_cell.value = label
    first_cell.fill = PatternFill("solid", fgColor=color)
    first_cell.font = Font(color="FFFFFF", bold=True, size=9)
    first_cell.alignment = Alignment(horizontal="center", vertical="center")

# En-tetes ligne 2
ws.row_dimensions[2].height = 50
for col_idx, (header, field, list_key, comment) in enumerate(COLUMNS, start=1):
    color = BLOC_COLORS.get(field, "555555")
    cell = ws.cell(row=2, column=col_idx, value=header)
    cell.fill      = PatternFill("solid", fgColor=color)
    cell.font      = Font(color="FFFFFF", bold=True, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border_normal
    cell.comment   = Comment(comment, "Guide")

# Cellules de donnees
for row in range(3, NB_ROWS + 3):
    ws.row_dimensions[row].height = 20
    for col_idx, (header, field, list_key, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx)
        cell.font      = Font(size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border    = border_normal
        cell.fill      = PatternFill("solid", fgColor=BLOC_CELL_COLORS.get(field, "FFFFFF"))

# Validations deroulantes
for col_idx, (header, field, list_key, _) in enumerate(COLUMNS, start=1):
    if list_key and list_key in list_col_map:
        col_letter, nb_vals = list_col_map[list_key]
        formula = "Listes!${}$2:${}${}".format(col_letter, col_letter, nb_vals + 1)
        dv = DataValidation(
            type="list", formula1=formula, allow_blank=True,
            showErrorMessage=True, errorTitle="Valeur invalide",
            error="Choisissez une valeur dans la liste deroulante.",
        )
        col_data = get_column_letter(col_idx)
        dv.sqref = "{}3:{}{}".format(col_data, col_data, NB_ROWS + 2)
        ws.add_data_validation(dv)

# Largeurs
WIDTHS = {
    "statut": 26, "type_habitation": 14, "type_residence": 18,
    "etage": 18, "surface": 12, "anciennete": 18, "jours_inhabita": 20,
    "eq_garage": 10, "eq_antivol": 14, "eq_veranda": 16, "surface_veranda": 16,
    "eq_annexes": 18, "surface_annexes": 16, "eq_chauffage": 18,
    "chauffage": 22, "cheminee_pro": 20, "eq_piscine": 14, "piscine": 26,
    "capital_mobilier": 18, "nb_adultes": 10, "nb_enfants": 10,
    "etat_assureur": 22, "assureur": 18,
    "date_naissance": 18, "situation": 18, "profession": 26,
    "cp": 12, "ville": 16,
}
for col_idx, (header, field, _, _) in enumerate(COLUMNS, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = WIDTHS.get(field, 14)

ws.freeze_panes = "A3"

# Legende
ws_leg = wb.create_sheet("Legende")
ws_leg.column_dimensions["A"].width = 30
ws_leg.column_dimensions["B"].width = 65

legend = [
    ("GUIDE DE REMPLISSAGE",    ""),
    ("", ""),
    ("Couleur",                 "Section"),
    ("Bleu fonce",              "Informations sur le logement"),
    ("Bleu clair",              "Equipements du logement (OUI / NON par equipement)"),
    ("Vert",                    "Occupants et capital mobilier"),
    ("Ocre / jaune",            "Assurance actuelle"),
    ("Violet",                  "Profil de l assure"),
    ("Marron",                  "Localisation"),
    ("", ""),
    ("REGLES IMPORTANTES",      ""),
    ("Equipements",             "Mettre OUI ou NON sur chaque equipement independamment - combinaisons libres"),
    ("Etage",                   "Remplir uniquement pour un appartement - laisser vide pour une maison"),
    ("Type chauffage",          "Remplir uniquement si Chauffage bois/cheminee = OUI"),
    ("Cheminee pro",            "Remplir uniquement si cheminee avec ou sans insert"),
    ("Type piscine",            "Remplir uniquement si Piscine = OUI"),
    ("Surface veranda/annexes", "Remplir uniquement si l equipement correspondant est OUI"),
    ("Assureur actuel",         "Remplir uniquement si Deja assure = Oui je suis assure"),
    ("Dates",                   "Format JJ/MM/AAAA  ex: 15/06/1985"),
    ("Surface",                 "En m2 entier  ex: 60"),
    ("Capital mobilier",        "En euros entier  ex: 15000"),
    ("Code postal",             "Laisser vide = Paris (75001) par defaut"),
    ("", ""),
    ("LANCER LE SCRAPING",      "python scraper.py --from-excel profils_template.xlsx --no-resume"),
]
for r_idx, (a, b) in enumerate(legend, start=1):
    ca = ws_leg.cell(row=r_idx, column=1, value=a)
    cb = ws_leg.cell(row=r_idx, column=2, value=b)
    if a in ("GUIDE DE REMPLISSAGE", "REGLES IMPORTANTES", "LANCER LE SCRAPING"):
        ca.font = Font(bold=True, size=11)
    elif a == "Couleur":
        ca.font = Font(bold=True)
        cb.font = Font(bold=True)

ws_lists.sheet_state = "hidden"
wb.save("profils_template.xlsx")
print("Fichier cree : profils_template.xlsx")
